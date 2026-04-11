import os
import re
import yaml
from docx import Document
from docx.shared import Pt
from pathlib import Path

def md_to_docx(md_path, docx_path):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Remove YAML frontmatter
    if content.startswith('---'):
        parts = content.split('---', 2)
        if len(parts) >= 3:
            content = parts[2].strip()

    # Clean characters incompatible with XML (Word requirement)
    # XML valid chars: #x9 | #xA | #xD | [#x20-#xD7FF] | [#xE000-#xFFFD] | [#x10000-#x10FFFF]
    def is_xml_valid(char):
        cp = ord(char)
        return (
            cp == 0x9 or cp == 0xA or cp == 0xD or
            (0x20 <= cp <= 0xD7FF) or
            (0xE000 <= cp <= 0xFFFD) or
            (0x10000 <= cp <= 0x10FFFF)
        )
    content = "".join(c for c in content if is_xml_valid(c))

    doc = Document()
    doc.add_heading('Converted Commentary Document', 0)
    
    # Simple line-by-line processing
    lines = content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph('')
            continue
            
        # Handle headers
        header_match = re.match(r'^(#+)\s+(.*)', line)
        if header_match:
            level = len(header_match.group(1))
            text = header_match.group(2)
            doc.add_heading(text, level=min(level, 4))
        else:
            # Add simple paragraph
            p = doc.add_paragraph(line)
            
    print(f"Saving to {docx_path}...")
    doc.save(docx_path)

def main():
    base_dir = Path("RAG_Output_Gemini/commentary")
    for subdir in base_dir.iterdir():
        if subdir.is_dir():
            # Nestled structure check
            md_file = subdir / subdir.name / f"{subdir.name}.md"
            if md_file.exists():
                docx_file = subdir / subdir.name / f"{subdir.name}.docx"
                md_to_docx(md_file, docx_file)

if __name__ == "__main__":
    main()
