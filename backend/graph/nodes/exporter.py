import os
import zipfile
import shutil
from typing import Any, Dict, List
import textwrap
from backend.graph.state import AgentState
from backend.services.bus import broadcast_progress
from backend.services.serialization import extract_text_content, ensure_blocks, make_stage, blocks_to_text

class ExporterNode:
    # ... (header same)
    def __init__(self):
        self.base_dir = os.path.abspath(os.getcwd())
        self.export_dir = os.path.join(self.base_dir, "backend", "static", "exports")
        self.plots_dir = os.path.join(self.base_dir, "backend", "static", "plots")
        os.makedirs(self.export_dir, exist_ok=True)

    def _create_latex_wrapper(self, content):
        """学术 LaTeX 模板封装"""
        template = r"""\documentclass{article}
\usepackage[utf8]{inputenc}
\usepackage{amsmath, amssymb}
\usepackage{graphicx}

\title{MCM/ICM Research Paper}
\author{MM-Agent Digital Orchestration}
\date{\today}

\begin{document}
\maketitle
"""
        template += content
        template += r"\n\end{document}"
        return template

    def _paper_blocks_to_markdown(self, blocks_like: Any) -> str:
        blocks = ensure_blocks(blocks_like)
        lines: List[str] = []
        for block in blocks:
            btype = block.get("type")
            if btype in {"text", "markdown"}:
                lines.append(str(block.get("text", "")))
            elif btype == "code":
                lang = str(block.get("language", "text"))
                code = str(block.get("code", ""))
                lines.append(f"```{lang}\n{code}\n```")
            elif btype == "math":
                latex = str(block.get("latex", ""))
                lines.append(f"$$\n{latex}\n$$")
            elif btype == "image":
                url = str(block.get("url", ""))
                alt = str(block.get("alt", "figure"))
                if url:
                    lines.append(f"![{alt}]({url})")
            elif btype == "table":
                headers = block.get("headers") or []
                rows = block.get("rows") or []
                if headers:
                    lines.append("| " + " | ".join(str(h) for h in headers) + " |")
                    lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for row in rows:
                    lines.append("| " + " | ".join(str(cell) for cell in (row or [])) + " |")
        return "\n\n".join([ln for ln in lines if ln.strip()]).strip()

    def _paper_blocks_to_plain_text(self, blocks_like: Any) -> str:
        return blocks_to_text(blocks_like) or ""

    def _generate_docx(self, blocks_like: Any, output_path: str) -> bool:
        try:
            from docx import Document  # type: ignore
        except Exception:
            return False
        try:
            blocks = ensure_blocks(blocks_like)
            doc = Document()
            doc.add_heading("MCM/ICM Research Paper", level=0)
            for block in blocks:
                btype = block.get("type")
                if btype in {"text", "markdown"}:
                    text = str(block.get("text", "")).strip()
                    if not text:
                        continue
                    for line in text.splitlines():
                        s = line.strip()
                        if s.startswith("## "):
                            doc.add_heading(s[3:].strip(), level=2)
                        elif s.startswith("# "):
                            doc.add_heading(s[2:].strip(), level=1)
                        elif s:
                            doc.add_paragraph(s)
                elif btype == "code":
                    doc.add_paragraph(str(block.get("code", "")))
                elif btype == "math":
                    doc.add_paragraph(str(block.get("latex", "")))
                elif btype == "image":
                    url = str(block.get("url", "")).strip()
                    if url.startswith("/plots/"):
                        abs_path = os.path.join(self.base_dir, "backend", "static", "plots", url.split("/plots/", 1)[1].replace("/", os.sep))
                        if os.path.exists(abs_path):
                            try:
                                doc.add_picture(abs_path)
                            except Exception:
                                doc.add_paragraph(f"[Figure] {url}")
                        else:
                            doc.add_paragraph(f"[Figure] {url}")
            doc.save(output_path)
            return True
        except Exception:
            return False

    def _generate_pdf(self, plain_text: str, output_path: str) -> bool:
        # 优先使用 PyMuPDF（fitz）生成基础 PDF，确保导出链路可用。
        try:
            import fitz  # type: ignore
        except Exception:
            return False
        try:
            doc = fitz.open()
            page = doc.new_page()
            y = 50
            max_width = 500
            line_height = 16
            for para in (plain_text or "").split("\n"):
                wrapped = textwrap.wrap(para, width=55) or [""]
                for line in wrapped:
                    if y > 780:
                        page = doc.new_page()
                        y = 50
                    page.insert_text((50, y), line, fontsize=11)
                    y += line_height
                y += 4
            doc.save(output_path)
            doc.close()
            return True
        except Exception:
            return False

    async def __call__(self, state: AgentState):
        """执行物理导出流程"""
        shared_mem = state.get("shared_memory", {})
        paper_blocks = ensure_blocks(shared_mem.get("paper_draft", "暂无内容。"))
        draft = self._paper_blocks_to_markdown(paper_blocks) or extract_text_content(shared_mem.get("paper_draft", "暂无内容。"))
        plain_text = self._paper_blocks_to_plain_text(paper_blocks) or draft
        artifacts_manifest = shared_mem.get("artifacts_manifest", [])
        
        # 1. 广播进度：开始文件转换
        await broadcast_progress("Exporter", "正在生成 Markdown 与 LaTeX 物理文件...", 20)
        
        md_path = os.path.join(self.export_dir, "thesis.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(draft)
        
        tex_path = os.path.join(self.export_dir, "thesis.tex")
        with open(tex_path, "w", encoding="utf-8") as f:
            f.write(self._create_latex_wrapper(draft))
        docx_path = os.path.join(self.export_dir, "thesis.docx")
        pdf_path = os.path.join(self.export_dir, "thesis.pdf")
        docx_ok = self._generate_docx(paper_blocks, docx_path)
        pdf_ok = self._generate_pdf(plain_text, pdf_path)
            
        # 2. 广播进度：开始全量打包
        await broadcast_progress("Exporter", "正在打包全量资产 (ZIP)...", 60)
        
        zip_path = os.path.join(self.export_dir, "final_submission.zip")
        with zipfile.ZipFile(zip_path, 'w') as zipf:
            zipf.write(md_path, arcname="thesis.md")
            zipf.write(tex_path, arcname="thesis.tex")
            if docx_ok and os.path.exists(docx_path):
                zipf.write(docx_path, arcname="thesis.docx")
            if pdf_ok and os.path.exists(pdf_path):
                zipf.write(pdf_path, arcname="thesis.pdf")
            
            # 添加动态产物清单中的文件
            if isinstance(artifacts_manifest, list):
                for item in artifacts_manifest:
                    if not isinstance(item, dict):
                        continue
                    abs_path = str(item.get("path", "")).strip()
                    filename = str(item.get("filename", "")).strip()
                    kind = str(item.get("kind", "")).strip().lower()
                    if not abs_path or not os.path.exists(abs_path):
                        continue
                    if kind == "image":
                        arcname = f"plots/{filename or os.path.basename(abs_path)}"
                    else:
                        arcname = f"exports/{filename or os.path.basename(abs_path)}"
                    zipf.write(abs_path, arcname=arcname)
            else:
                # 兼容兜底
                plot_file = os.path.join(self.plots_dir, "output.png")
                if os.path.exists(plot_file):
                    zipf.write(plot_file, arcname="plots/output.png")
            
        await broadcast_progress("Exporter", "资产固化完成，ZIP 已就绪。", 100)
        
        report_lines = ["✅ 成果物理导出成功！您可以直接下载以下资源："]
        if pdf_ok and os.path.exists(pdf_path):
            report_lines.append("- [PDF 版本](/exports/thesis.pdf)")
        if docx_ok and os.path.exists(docx_path):
            report_lines.append("- [DOCX 版本](/exports/thesis.docx)")
        report_lines.extend(
            [
                "- [Markdown 论文原文](/exports/thesis.md)",
                "- [LaTeX 专业版本](/exports/thesis.tex)",
                "- [全量压缩包 (含图表)](/exports/final_submission.zip)",
            ]
        )
        report = "\n".join(report_lines) + "\n"

        return {
            "status": "APPROVED",
            "messages": [{"role": "ai", "content": ensure_blocks(report)}],
            "draft": ensure_blocks(report),
            "stage": make_stage("export_completed", "Physical Asset固化完成"),
        }

exporter_node = ExporterNode()
